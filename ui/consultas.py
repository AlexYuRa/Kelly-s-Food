# -*- coding: utf-8 -*-
"""
Kelly's Food - Consultas y Reportes
===================================
Consultas de negocio predefinidas con filtro de fechas, resultado en tabla y
exportación a CSV o Excel.
"""

import csv
from tkinter import ttk, filedialog, messagebox

import customtkinter as ctk

import db
from ui import estilos


# Cada reporte: nombre visible -> función(desde, hasta) que devuelve
# (encabezados, filas) donde filas es lista de tuplas.
def _rep_ingresos_gastos(desde, hasta):
    filas = db.ingresos_gastos_por_mes(desde, hasta)
    datos = [(f["mes"], round(f["ingreso"] or 0, 2), round(f["gasto"] or 0, 2),
              round((f["ingreso"] or 0) - (f["gasto"] or 0), 2)) for f in filas]
    return ["Mes", "Ingresos", "Gastos", "Balance"], datos


def _rep_raciones_mes(desde, hasta):
    filas = db.raciones_por_mes(desde, hasta)
    return ["Mes", "Raciones servidas"], [(f["mes"], f["total"]) for f in filas]


def _rep_top_raciones(desde, hasta):
    filas = db.top_trabajadores_raciones(20, desde, hasta)
    return (["ID Trabajador", "Nombre", "Raciones"],
            [(f["id_trabajador"], f["nombre"], f["total"]) for f in filas])


def _rep_top_pagos(desde, hasta):
    filas = db.top_trabajadores_pagos(20, desde, hasta)
    return (["ID Trabajador", "Nombre", "Total pagado (S/)"],
            [(f["id_trabajador"], f["nombre"], round(f["total"], 2)) for f in filas])


def _rep_metodo_pago(desde, hasta):
    filas = db.uso_metodo_pago(desde, hasta)
    return (["Método de pago", "Nº pagos", "Total (S/)"],
            [(f["Metodo_Pago"], f["n"], round(f["total"], 2)) for f in filas])


REPORTES = {
    "Ingresos vs gastos por mes": _rep_ingresos_gastos,
    "Raciones servidas por mes": _rep_raciones_mes,
    "Top 20 trabajadores por raciones": _rep_top_raciones,
    "Top 20 trabajadores por monto pagado": _rep_top_pagos,
    "Uso de métodos de pago": _rep_metodo_pago,
}


class Consultas(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color="transparent")
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)
        self._encabezados = []
        self._filas = []
        self._lo, self._hi = db.rango_fechas()
        self._construir_controles()
        self._construir_tabla()
        self.ejecutar()

    def _construir_controles(self):
        ctk.CTkLabel(self, text="Consultas y Reportes",
                     font=estilos.fuente(24, "bold")).grid(
            row=0, column=0, sticky="w", padx=4, pady=(4, 10))

        barra = ctk.CTkFrame(self, corner_radius=12)
        barra.grid(row=1, column=0, sticky="ew", padx=4, pady=(0, 10))
        barra.grid_columnconfigure(0, weight=1)

        # Fila 1: selección de reporte + filtros de fecha + Ejecutar.
        fila1 = ctk.CTkFrame(barra, fg_color="transparent")
        fila1.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 4))

        ctk.CTkLabel(fila1, text="Reporte", font=estilos.fuente(12, "bold")).pack(
            side="left", padx=(0, 6))
        self.combo = ctk.CTkComboBox(fila1, values=list(REPORTES.keys()),
                                     width=260, height=34, command=lambda _=None: self.ejecutar())
        self.combo.set(next(iter(REPORTES)))
        self.combo.pack(side="left", padx=(0, 16))

        ctk.CTkLabel(fila1, text="Desde", font=estilos.fuente(12)).pack(side="left", padx=(0, 4))
        self.e_desde = ctk.CTkEntry(fila1, width=105, height=34)
        self.e_desde.insert(0, self._lo or "")
        self.e_desde.pack(side="left", padx=(0, 10))
        ctk.CTkLabel(fila1, text="Hasta", font=estilos.fuente(12)).pack(side="left", padx=(0, 4))
        self.e_hasta = ctk.CTkEntry(fila1, width=105, height=34)
        self.e_hasta.insert(0, self._hi or "")
        self.e_hasta.pack(side="left", padx=(0, 10))
        ctk.CTkButton(fila1, text="Ejecutar", width=100, height=34,
                      command=self.ejecutar).pack(side="left", padx=6)

        # Fila 2: acciones de exportación (alineadas a la derecha).
        fila2 = ctk.CTkFrame(barra, fg_color="transparent")
        fila2.grid(row=1, column=0, sticky="ew", padx=12, pady=(0, 12))
        ctk.CTkButton(fila2, text="⬇ Exportar Excel", width=150, height=34,
                      fg_color=estilos.AMBAR, hover_color=estilos.AMBAR_HOVER,
                      text_color="#20302A", command=self.exportar_excel).pack(side="right")
        ctk.CTkButton(fila2, text="⬇ Exportar CSV", width=140, height=34,
                      fg_color=estilos.AMBAR, hover_color=estilos.AMBAR_HOVER,
                      text_color="#20302A", command=self.exportar_csv).pack(side="right", padx=(0, 8))

    def _construir_tabla(self):
        cont = ctk.CTkFrame(self, corner_radius=12)
        cont.grid(row=2, column=0, sticky="nsew", padx=4, pady=(0, 4))
        cont.grid_rowconfigure(0, weight=1)
        cont.grid_columnconfigure(0, weight=1)

        estilos.estilizar_treeview()
        self.tree = ttk.Treeview(cont, show="headings", style="Kelly.Treeview")
        vsb = ttk.Scrollbar(cont, orient="vertical", command=self.tree.yview,
                            style="Kelly.Vertical.TScrollbar")
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew", padx=(6, 0), pady=6)
        vsb.grid(row=0, column=1, sticky="ns", pady=6, padx=(0, 6))

        self.lbl = ctk.CTkLabel(cont, text="", font=estilos.fuente(11),
                                text_color=estilos.TEXTO_TENUE)
        self.lbl.grid(row=1, column=0, sticky="w", padx=10, pady=(0, 6))

    def ejecutar(self):
        desde = self.e_desde.get().strip() or None
        hasta = self.e_hasta.get().strip() or None
        try:
            self._encabezados, self._filas = REPORTES[self.combo.get()](desde, hasta)
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo ejecutar la consulta:\n{e}")
            return

        self.tree["columns"] = self._encabezados
        self.tree.delete(*self.tree.get_children())
        for h in self._encabezados:
            self.tree.heading(h, text=h)
            self.tree.column(h, width=160, anchor="w")
        self.tree.tag_configure("par", background=estilos._lado(
            estilos.SUPERFICIE, ctk.get_appearance_mode()))
        for i, fila in enumerate(self._filas):
            self.tree.insert("", "end", values=fila, tags=("par",) if i % 2 else ())
        self.lbl.configure(text=f"{len(self._filas)} fila(s)")

    # ------------------------------------------------------------------
    def _sin_datos(self):
        if not self._filas:
            messagebox.showinfo("Sin datos", "No hay filas para exportar.")
            return True
        return False

    def exportar_csv(self):
        if self._sin_datos():
            return
        ruta = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV", "*.csv")],
            initialfile=self._nombre_archivo())
        if not ruta:
            return
        try:
            with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(self._encabezados)
                w.writerows(self._filas)
            messagebox.showinfo("Exportado", f"CSV guardado en:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo exportar:\n{e}")

    def exportar_excel(self):
        if self._sin_datos():
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            messagebox.showerror(
                "Falta openpyxl",
                "Instala la dependencia:\n    py -m pip install openpyxl")
            return
        ruta = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=self._nombre_archivo())
        if not ruta:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte"
            ws.append(self._encabezados)
            for celda in ws[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="2E7D5B")
            for fila in self._filas:
                ws.append(list(fila))
            for col in ws.columns:
                ancho = max(len(str(c.value)) for c in col) + 2
                ws.column_dimensions[col[0].column_letter].width = min(ancho, 40)
            wb.save(ruta)
            messagebox.showinfo("Exportado", f"Excel guardado en:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo exportar:\n{e}")

    def _nombre_archivo(self):
        base = self.combo.get().lower().replace(" ", "_")
        base = "".join(ch for ch in base if ch.isalnum() or ch == "_")
        return f"reporte_{base}"
